"""
io_helpers.py - 결과 저장 (CSV / JSON / Excel)

save_variable_csvs() : 모든 Variable → results/*.csv
save_results_json()  : 구조화 결과 → results/summary.json (재현성)
export_summary_excel(): 요약·설치·비용 분해 엑셀
"""

import json
import logging
from pathlib import Path

from src.utils.config_loader import get

logger = logging.getLogger("cired")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    _HAS_XLSX = True
except ImportError:
    _HAS_XLSX = False


def _sanitize(obj):
    """numpy/int-key → JSON 직렬화 가능 형태."""
    if isinstance(obj, dict):
        return {str(k): _sanitize(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_sanitize(v) for v in obj]
    if hasattr(obj, "item"):              # numpy scalar
        return obj.item()
    return obj


def save_variable_csvs(results_dfs: dict, out_dir):
    """extract_all() 결과 dict({name: DataFrame}) 를 CSV 로 저장."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, df in results_dfs.items():
        safe = name.replace("[", "").replace("]", "").replace(".", "_")
        df.to_csv(out_dir / f"{safe}.csv", index=False)
    logger.info(f"변수 CSV {len(results_dfs)}개 저장 → {out_dir}")


def save_results_json(results: dict, out_dir):
    """구조화 결과를 JSON 으로 저장."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "summary.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(_sanitize(results), f, ensure_ascii=False, indent=2)
    logger.info(f"요약 JSON 저장 → {path}")


# ─────────────────────────────────────────────────────────
# Excel
# ─────────────────────────────────────────────────────────
_BLUE = "1F4E79"
_GREEN = "1E5631"
_YELLOW = "FFEB9C"


def _hdr(ws, row, headers, fill):
    border = Border(*[Side(style="thin")] * 4)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = PatternFill(fill_type="solid", fgColor=fill)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border


def _row(ws, row, values, fmt=None):
    border = Border(*[Side(style="thin")] * 4)
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.alignment = Alignment(horizontal="center")
        c.border = border
        if fmt and isinstance(v, float):
            c.number_format = fmt


def export_summary_excel(results: dict, config: dict, path):
    """요약 + ESS/CHP 설치 + 비용 분해 엑셀."""
    if not _HAS_XLSX:
        logger.warning("openpyxl 없음 — 엑셀 저장 건너뜀")
        return
    wb = Workbook()

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"
    for col, w in enumerate([26, 18, 12], 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    _hdr(ws, 1, ["항목", "값", "단위"], _BLUE)
    rows = [
        ("계획 전략", results.get("mode", ""), ""),
        ("총비용 (목적함수)", round(results.get("objective", 0) / 1e8, 2), "억원"),
        ("투자비", round(results.get("invest_cost", 0) / 1e8, 2), "억원"),
        ("기대 운영비", round(results.get("expected_op_cost", 0) / 1e8, 2), "억원"),
        ("CVaR 운영비", round(results.get("cvar_op_cost", 0) / 1e8, 2), "억원"),
        ("PV 총용량", results.get("total_PV_MWp", 0), "MWp"),
        ("ESS 총출력", results.get("total_P_MW", 0), "MW"),
        ("ESS 총에너지", results.get("total_E_MWh", 0), "MWh"),
        ("CHP 총전기용량", results.get("total_CHP_MW", 0), "MW"),
        ("PV curtailment율", results.get("curtailment_pct", 0), "%"),
        ("REC 수익", results.get("rec_revenue_억", 0), "억원"),
        ("ESS 설치 노드 수", results.get("n_ess_sites", 0), "개소"),
        ("PV 설치 노드 수", results.get("n_pv_sites", 0), "개소"),
        ("CHP 설치 노드 수", results.get("n_chp_sites", 0), "개소"),
    ]
    for i, r in enumerate(rows, 2):
        _row(ws, i, list(r))

    # Sheet 2: ESS_Siting
    ws2 = wb.create_sheet("ESS_Siting")
    _hdr(ws2, 1, ["버스", "P [MW]", "E [MWh]", "P:E [h]"], _BLUE)
    for i, (n, v) in enumerate(sorted(results.get("ess_sites", {}).items()), 2):
        pe = v["E_MWh"] / v["P_MW"] if v["P_MW"] > 0 else 0
        _row(ws2, i, [f"Bus {int(n) + 1}", v["P_MW"], v["E_MWh"], round(pe, 2)])

    # Sheet 3: CHP_Siting
    if results.get("chp_sites"):
        ws3 = wb.create_sheet("CHP_Siting")
        _hdr(ws3, 1, ["버스", "P_e [MW]", "P_h [MW]", "보일러 [MW]", "TES [MWh]"], _GREEN)
        for i, (n, v) in enumerate(sorted(results["chp_sites"].items()), 2):
            _row(ws3, i, [f"Bus {int(n) + 1}", v["P_e_MW"], v["P_h_MW"],
                          v["P_boiler_MW"], v["E_tes_MWh"]])

    # Sheet 4: Cost_Breakdown
    ws4 = wb.create_sheet("Cost_Breakdown")
    _hdr(ws4, 1, ["항목", "값 [억원]"], _YELLOW)
    cb = [
        ("투자비", round(results.get("invest_cost", 0) / 1e8, 2)),
        ("기대 운영비", round(results.get("expected_op_cost", 0) / 1e8, 2)),
        ("CVaR 운영비", round(results.get("cvar_op_cost", 0) / 1e8, 2)),
        ("REC 수익", -results.get("rec_revenue_억", 0)),
        ("총비용(목적함수)", round(results.get("objective", 0) / 1e8, 2)),
    ]
    for i, r in enumerate(cb, 2):
        _row(ws4, i, list(r))

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info(f"요약 엑셀 저장 → {path}")
